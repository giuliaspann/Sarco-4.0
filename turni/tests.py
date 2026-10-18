from datetime import date, datetime, timedelta
from io import BytesIO, StringIO
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from . import importa
from .models import Dipendente, ElementoPulizia, Fascia, Postazione, Pulizia, RegolaPulizia, Turno

LUNEDI = date(2026, 10, 12)


def file_turni(righe, note=(), foglio='Sett.42'):
    """Un Excel con la stessa forma di quello vero."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = foglio
    ws['B4'] = 'DIPENDENTI'
    for i, g in enumerate(["LUNEDI'", "MARTEDI'", "MERCOLEDI'", "GIOVEDI'", "VENERDI'", 'SABATO', 'DOMENICA']):
        ws.cell(4, 3 + i, g)
    for r, (nome, *turni) in enumerate(righe, start=5):
        ws.cell(r, 2, nome)
        for c, t in enumerate(turni, start=3):
            ws.cell(r, c, t)
    for r, (rif, dip, testo) in enumerate(note, start=10):
        ws.cell(r, 11, rif)
        ws.cell(r, 12, dip)
        ws.cell(r, 13, testo)
    ws.cell(30, 8, 'APPRENDISTI E TUTOR')
    ws.cell(31, 8, "DRAME BUOBA (LAUDICINA E SPANO')")
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'DAL 12-10-2026 AL 18-10-2026.xlsx'
    return buf


RIGHE = [
    ("SPANO' A. MAURO", '1T- Imp. 1 ', 'RIP', 'CAPO T. - Imp. 1', '1T- Imp. 1 ', 'FERIE', 'RIP', 'FESTA'),
    ('DRAME BOUBA', '2T- Imp. 1 ', '2T-Imp. 2', 'RIP', '1T- Imp. 1 ', '2T- Imp. 1 ', 'C.P.', 'RIP'),
    ('CRIMI BENITO', 'N/palista', 'CAPO T. Imp. 1 N', 'N/imp.2', 'N', 'G-Analisi', 'G - Mulino', '1T/PALISTA'),
]


class CodiciTest(TestCase):
    def test_codici_del_file(self):
        casi = {
            '1T- Imp. 1 ': (Fascia.MATTINA, Postazione.IMP1, False),
            '2T-Imp. 2': (Fascia.POMERIGGIO, Postazione.IMP2, False),
            'CAPO T. - Imp. 1': (Fascia.ALTRO, Postazione.IMP1, True),
            'CAPO T.- Imp. 1 ': (Fascia.ALTRO, Postazione.IMP1, True),
            'CAPO T. Imp. 1 N': (Fascia.NOTTE, Postazione.IMP1, True),
            'N/palista': (Fascia.NOTTE, Postazione.PALISTA, False),
            'N/imp.2': (Fascia.NOTTE, Postazione.IMP2, False),
            'N': (Fascia.NOTTE, '', False),
            'G': (Fascia.GIORNALIERO, '', False),
            'G-Analisi': (Fascia.GIORNALIERO, Postazione.ANALISI, False),
            '1T-Analisi': (Fascia.MATTINA, Postazione.ANALISI, False),
            'G - Mulino': (Fascia.GIORNALIERO, Postazione.MULINO, False),
            'G-BARROTTI': (Fascia.GIORNALIERO, Postazione.BARROTTI, False),
            '1T-Superbi': (Fascia.MATTINA, Postazione.SUPERBI, False),
            '2T/PALISTA': (Fascia.POMERIGGIO, Postazione.PALISTA, False),
            'RIP': (Fascia.RIPOSO, '', False),
            'FERIE': (Fascia.FERIE, '', False),
            'FESTA': (Fascia.FESTA, '', False),
            'C.P.': (Fascia.PERMESSO, '', False),
            None: (Fascia.RIPOSO, '', False),
        }
        for codice, atteso in casi.items():
            c = importa.interpreta_codice(codice)
            self.assertEqual((c.fascia, c.postazione, c.capoturno), atteso, codice)

    def test_nomi_scritti_diversi(self):
        self.assertTrue(importa.stesso_nome("SPANO' A. MAURO", 'SPANO MAURO'))
        self.assertTrue(importa.stesso_nome('SAMMARTANO CARLO', 'CARLO SAMMARTANO'))
        self.assertTrue(importa.stesso_nome('DRAME BOUBA', 'DRAME BUOBA'))
        self.assertFalse(importa.stesso_nome('LAUDICINA DAVIDE', 'LAUDICINA GIUSEPPE'))
        self.assertFalse(importa.stesso_nome('DI GIROLAMO GIANVITO', 'DI GIROLAMO FRANCESCO'))
        self.assertFalse(importa.stesso_nome('ALAGNA PIETRO', 'ALAGNA GIUSEPPE'))

    def test_lunedi_dal_nome_file(self):
        self.assertEqual(importa.lunedi_da_nome_file('DAL 12-10-2026 AL 18-10-2026.xlsx'), LUNEDI)
        self.assertEqual(importa.lunedi_da_nome_file('DAL 14-10-2026.xlsx'), LUNEDI)
        self.assertIsNone(importa.lunedi_da_nome_file('turni.xlsx'))


class ImportTest(TestCase):
    def test_lettura_e_salvataggio(self):
        note = [('42', 'SPANO MAURO', 'ferie venerdì'), ('43', 'CRIMI BENITO', 'fare notte')]
        letta = importa.leggi_file(file_turni(RIGHE, note), LUNEDI)
        self.assertEqual(letta.foglio, 'Sett.42')
        self.assertEqual(len(letta.righe), 3)
        settimana, nuovi, _ = importa.salva_settimana(letta)
        self.assertEqual(len(nuovi), 3)
        self.assertEqual(Turno.objects.count(), 21)
        crimi = Turno.objects.get(dipendente__nome='CRIMI BENITO', data=LUNEDI + timedelta(days=1))
        self.assertTrue(crimi.capoturno)
        self.assertEqual(crimi.fascia, Fascia.NOTTE)
        # il tutor si aggancia anche se il nome è scritto diverso
        self.assertEqual(Dipendente.objects.get(nome='DRAME BOUBA').tutor, "LAUDICINA E SPANO'")
        self.assertEqual(settimana.note.count(), 2)

    def test_ricaricare_non_perde_le_pulizie(self):
        importa.salva_settimana(importa.leggi_file(file_turni(RIGHE), LUNEDI))
        turno = Turno.objects.get(dipendente__nome="SPANO' A. MAURO", data=LUNEDI)
        el = ElementoPulizia.objects.create(nome='Nastro', tipo='ZONA')
        Pulizia.objects.create(turno=turno, elemento=el, stato=Pulizia.Stato.FATTA)

        corrette = [list(r) for r in RIGHE]
        corrette[0][1] = '2T- Imp. 1 '
        importa.salva_settimana(importa.leggi_file(file_turni(corrette), LUNEDI))
        turno.refresh_from_db()
        self.assertEqual(turno.fascia, Fascia.POMERIGGIO)
        self.assertEqual(turno.pulizie.count(), 1)
        self.assertEqual(Dipendente.objects.count(), 3)

    def test_foglio_mancante(self):
        letta = importa.leggi_file(file_turni(RIGHE, foglio='Foglio1'), LUNEDI)
        self.assertEqual(letta.foglio, 'Foglio1')
        self.assertTrue(letta.avvisi)


class PulizieTest(TestCase):
    def setUp(self):
        importa.salva_settimana(importa.leggi_file(file_turni(RIGHE), LUNEDI))
        self.nastro = ElementoPulizia.objects.create(nome='Nastro Imp. 1', tipo='ZONA')
        self.pala = ElementoPulizia.objects.create(nome='Pala gommata', tipo='MEZZO')
        self.ufficio = ElementoPulizia.objects.create(nome='Sala controllo', tipo='ZONA')
        RegolaPulizia.objects.create(elemento=self.nastro, postazione=Postazione.IMP1)
        RegolaPulizia.objects.create(elemento=self.pala, postazione=Postazione.PALISTA)
        RegolaPulizia.objects.create(elemento=self.ufficio, solo_capoturno=True)

        User = get_user_model()
        self.spano = Dipendente.objects.get(nome="SPANO' A. MAURO")
        self.spano.utente = User.objects.create_user('spano.mauro', password='123456')
        self.spano.save()
        self.drame = Dipendente.objects.get(nome='DRAME BOUBA')
        self.drame.utente = User.objects.create_user('drame.bouba', password='654321')
        self.drame.save()

    def a(self, giorno, ora=12):
        return mock.patch(
            'django.utils.timezone.now',
            return_value=timezone.make_aware(datetime.combine(giorno, datetime.min.time()) + timedelta(hours=ora)),
        )

    def test_regole_per_ruolo(self):
        from .pulizie import genera_pulizie

        capo = Turno.objects.get(dipendente=self.spano, data=LUNEDI + timedelta(days=2))
        genera_pulizie(capo)
        self.assertEqual({p.elemento for p in capo.pulizie.all()}, {self.nastro, self.ufficio})
        palista = Turno.objects.get(dipendente__nome='CRIMI BENITO', data=LUNEDI)
        genera_pulizie(palista)
        self.assertEqual([p.elemento for p in palista.pulizie.all()], [self.pala])
        riposo = Turno.objects.get(dipendente=self.spano, data=LUNEDI + timedelta(days=1))
        genera_pulizie(riposo)
        self.assertFalse(riposo.pulizie.exists())

    def test_operaio_segna_fatta_e_capoturno_verifica(self):
        self.client.login(username='drame.bouba', password='654321')
        with self.a(LUNEDI + timedelta(days=3)):  # giovedì: Drame 1T Imp. 1
            r = self.client.get(reverse('turni:pulizie'))
            self.assertContains(r, 'Nastro Imp. 1')
            p = Pulizia.objects.get(turno__dipendente=self.drame)
            self.client.post(reverse('turni:pulizie_turno', args=[p.turno_id]), {'pulizia': p.pk, 'azione': 'fatta', 'nota': 'ok'})
            p.refresh_from_db()
            self.assertEqual(p.stato, Pulizia.Stato.FATTA)

            # Drame non è capoturno: niente pagina verifica
            self.assertEqual(self.client.get(reverse('turni:capoturno')).status_code, 403)

            # Spano giovedì è 1T, ma mercoledì era capoturno: vale anche il giorno dopo (turni di notte)
            self.client.login(username='spano.mauro', password='123456')
            r = self.client.get(reverse('turni:capoturno'))
            self.assertEqual(r.status_code, 200)
            self.client.post(reverse('turni:capoturno') + f'?giorno={p.turno.data}', {'pulizia': p.pk, 'azione': 'segnala', 'nota': 'rifare angolo'})
            p.refresh_from_db()
            self.assertEqual(p.stato, Pulizia.Stato.SEGNALATA)

    def test_non_si_toccano_turni_altrui(self):
        altrui = Turno.objects.get(dipendente=self.spano, data=LUNEDI)
        self.client.login(username='drame.bouba', password='654321')
        self.assertEqual(self.client.get(reverse('turni:pulizie_turno', args=[altrui.pk])).status_code, 404)

    def test_mezzo_aggiunto_dall_operaio(self):
        self.client.login(username='drame.bouba', password='654321')
        turno = Turno.objects.get(dipendente=self.drame, data=LUNEDI)
        url = reverse('turni:pulizie_turno', args=[turno.pk])
        self.client.post(url, {'azione': 'aggiungi', 'elemento': self.pala.pk})
        self.assertTrue(turno.pulizie.filter(elemento=self.pala, aggiunta_da_operaio=True).exists())

    def test_foto_ridotta_e_protetta(self):
        from PIL import Image

        img = BytesIO()
        Image.new('RGB', (4000, 3000), 'red').save(img, 'JPEG')
        img.seek(0)
        img.name = 'x.jpg'
        turno = Turno.objects.get(dipendente=self.drame, data=LUNEDI)
        self.client.login(username='drame.bouba', password='654321')
        with self.settings(MEDIA_ROOT=self._tmp()):
            self.client.get(reverse('turni:pulizie_turno', args=[turno.pk]))
            p = turno.pulizie.get()
            self.client.post(reverse('turni:pulizie_turno', args=[turno.pk]), {'pulizia': p.pk, 'azione': 'fatta', 'foto': img})
            p.refresh_from_db()
            self.assertEqual(Image.open(p.foto.path).size, (1600, 1200))
            self.assertEqual(self.client.get(reverse('turni:foto', args=[p.pk])).status_code, 200)
            # chi non è proprietario né capoturno non la vede
            User = get_user_model()
            crimi = Dipendente.objects.get(nome='CRIMI BENITO')
            crimi.utente = User.objects.create_user('crimi', password='111111')
            crimi.save()
            Turno.objects.filter(dipendente=crimi).update(capoturno=False)
            self.client.login(username='crimi', password='111111')
            with self.a(LUNEDI + timedelta(days=20)):
                self.assertEqual(self.client.get(reverse('turni:foto', args=[p.pk])).status_code, 403)

    def _tmp(self):
        import tempfile

        d = tempfile.mkdtemp()
        self.addCleanup(__import__('shutil').rmtree, d, True)
        return d


class AccessoTest(TestCase):
    def setUp(self):
        self.dip = Dipendente.objects.create(nome='ROSSI MARIO')
        self.dip.utente = get_user_model().objects.create_user('rossi.mario', password='123456')
        self.dip.save()

    def test_blocco_dopo_troppi_pin_sbagliati(self):
        url = reverse('turni:login')
        for _ in range(5):
            self.client.post(url, {'username': 'rossi.mario', 'password': '000000'})
        r = self.client.post(url, {'username': 'Rossi.Mario', 'password': '123456'})
        self.assertContains(r, 'Troppi tentativi')
        self.dip.refresh_from_db()
        self.dip.bloccato_fino = timezone.now() - timedelta(seconds=1)
        self.dip.save()
        r = self.client.post(url, {'username': 'Rossi.Mario', 'password': '123456'})
        self.assertRedirects(r, reverse('turni:home'), fetch_redirect_response=False)

    def test_pagine_ufficio_solo_staff(self):
        self.client.login(username='rossi.mario', password='123456')
        self.assertNotEqual(self.client.get(reverse('turni:importa')).status_code, 200)
        self.assertNotEqual(self.client.get(reverse('turni:dipendenti')).status_code, 200)

    def test_ufficio_carica_e_crea_pin(self):
        User = get_user_model()
        User.objects.create_user('ufficio', password='x', is_staff=True)
        self.client.login(username='ufficio', password='x')
        r = self.client.post(reverse('turni:importa'), {'file': file_turni(RIGHE)})
        self.assertRedirects(r, reverse('turni:anteprima'))
        self.assertContains(self.client.get(reverse('turni:anteprima')), 'CRIMI BENITO')
        self.client.post(reverse('turni:anteprima'))
        self.assertEqual(Turno.objects.count(), 21)

        crimi = Dipendente.objects.get(nome='CRIMI BENITO')
        r = self.client.post(reverse('turni:dipendenti'), {'dipendente': crimi.pk, 'azione': 'pin'}, follow=True)
        self.assertContains(r, 'crimi.benito')


class DeployTest(TestCase):
    def test_crea_utente_ufficio_una_volta(self):
        from django.core.management import call_command

        with mock.patch.dict('os.environ', {'SARCO_ADMIN_USER': 'ufficio', 'SARCO_ADMIN_PASSWORD': 'segreta-1'}):
            call_command('prepara_deploy', stdout=StringIO())
            u = get_user_model().objects.get(username='ufficio')
            self.assertTrue(u.is_superuser and u.check_password('segreta-1'))
            u.set_password('cambiata')
            u.save()
            call_command('prepara_deploy', stdout=StringIO())
            u.refresh_from_db()
            self.assertTrue(u.check_password('cambiata'))


class CalendarioTest(TestCase):
    def setUp(self):
        importa.salva_settimana(importa.leggi_file(file_turni(RIGHE), LUNEDI))
        d = Dipendente.objects.get(nome='CRIMI BENITO')
        d.utente = get_user_model().objects.create_user('crimi', password='123456')
        d.save()
        self.client.login(username='crimi', password='123456')

    def test_due_settimane_con_sigle_e_dettaglio_del_giorno(self):
        r = self.client.get(reverse('turni:miei_turni'), {'giorno': LUNEDI + timedelta(days=1)})
        self.assertEqual(r.status_code, 200)
        self.assertEqual([s['giorni'][0]['data'] for s in r.context['settimane']], [LUNEDI, LUNEDI + timedelta(days=7)])
        sigle = [c['turno'].sigla for s in r.context['settimane'] for c in s['giorni'] if c['turno']]
        self.assertEqual(sigle, ['N', 'N', 'N', 'N', 'G', 'G', '1T'])
        posti = [c['turno'].sigla_postazione for c in r.context['settimane'][0]['giorni']]
        self.assertEqual(posti, ['PALA', 'IMP 1', 'IMP 2', '', 'ANAL.', 'MUL.', 'PALA'])
        self.assertEqual(r.context['turno'].data, LUNEDI + timedelta(days=1))
        self.assertTrue(r.context['turno'].capoturno)

    def test_frecce_spostano_di_una_settimana(self):
        r = self.client.get(reverse('turni:miei_turni'), {'settimana': LUNEDI + timedelta(days=7)})
        self.assertEqual(r.context['inizio'], LUNEDI + timedelta(days=7))
        self.assertEqual(r.context['prec'], LUNEDI)
        self.assertIsNone(r.context['turno'])
        # toccare un giorno della seconda riga non sposta la finestra
        r = self.client.get(
            reverse('turni:miei_turni'), {'settimana': LUNEDI, 'giorno': LUNEDI + timedelta(days=9)}
        )
        self.assertEqual(r.context['inizio'], LUNEDI)
        self.assertEqual(r.context['scelto'], LUNEDI + timedelta(days=9))


class HomeTest(TestCase):
    def setUp(self):
        importa.salva_settimana(importa.leggi_file(file_turni(RIGHE), LUNEDI))
        for nome, user in [('CRIMI BENITO', 'crimi'), ('DRAME BOUBA', 'drame')]:
            d = Dipendente.objects.get(nome=nome)
            d.utente = get_user_model().objects.create_user(user, password='123456')
            d.save()
        adesso = timezone.make_aware(datetime.combine(LUNEDI + timedelta(days=1), datetime.min.time()) + timedelta(hours=12))
        patcher = mock.patch('django.utils.timezone.now', return_value=adesso)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_scorciatoie_operaio(self):
        self.client.login(username='drame', password='123456')
        r = self.client.get(reverse('turni:home'))
        self.assertEqual(r.status_code, 200)
        for testo in ('I miei turni', 'Pulizie', 'Cambio turno', 'Ferie / permesso', 'Avviso assenza', 'Chat', 'Bacheca'):
            self.assertContains(r, testo)
        self.assertEqual(r.context['turno'].data, LUNEDI + timedelta(days=1))  # martedì: 2T Imp. 2
        self.assertEqual(r.context['domani'].fascia, Fascia.RIPOSO)
        self.assertNotContains(r, 'Verifica pulizie')  # non è capo turno

    def test_capo_turno_vede_la_verifica(self):
        self.client.login(username='crimi', password='123456')  # martedì capo turno di notte
        self.assertContains(self.client.get(reverse('turni:home')), 'Verifica pulizie')


class GestioneDipendentiTest(TestCase):
    def setUp(self):
        get_user_model().objects.create_user('ufficio', password='x', is_staff=True)
        self.client.login(username='ufficio', password='x')

    def test_aggiungi_crea_accesso_ed_elimina(self):
        url = reverse('turni:dipendenti')
        self.client.post(url, {'azione': 'aggiungi', 'nome': '  spanò   giulia '})
        dip = Dipendente.objects.get(nome='SPANÒ GIULIA')
        r = self.client.post(url, {'azione': 'aggiungi', 'nome': 'Spanò Giulia'}, follow=True)
        self.assertContains(r, "già nell")
        self.assertEqual(Dipendente.objects.count(), 1)

        self.client.post(url, {'dipendente': dip.pk, 'azione': 'pin'})
        dip.refresh_from_db()
        utente = dip.utente
        self.assertIsNotNone(utente)

        self.client.post(url, {'dipendente': dip.pk, 'azione': 'elimina'})
        self.assertFalse(Dipendente.objects.exists())
        self.assertFalse(get_user_model().objects.filter(pk=utente.pk).exists())
        # l'utente dell'ufficio resta
        self.assertTrue(get_user_model().objects.filter(username='ufficio').exists())

    def test_operaio_non_gestisce_dipendenti(self):
        d = Dipendente.objects.create(nome='ROSSI MARIO')
        d.utente = get_user_model().objects.create_user('rossi', password='123456')
        d.save()
        self.client.login(username='rossi', password='123456')
        self.client.post(reverse('turni:dipendenti'), {'dipendente': d.pk, 'azione': 'elimina'})
        self.assertTrue(Dipendente.objects.filter(pk=d.pk).exists())


class TurniProvaTest(TestCase):
    def test_turni_di_prova_solo_per_chi_non_ne_ha(self):
        get_user_model().objects.create_user('ufficio', password='x', is_staff=True)
        self.client.login(username='ufficio', password='x')
        dip = Dipendente.objects.create(nome='SPANÒ GIULIA')
        url = reverse('turni:dipendenti')
        self.assertContains(self.client.get(url), 'Genera turni di prova')
        self.client.post(url, {'dipendente': dip.pk, 'azione': 'turni_prova'})
        turni = list(dip.turni.order_by('data'))
        self.assertEqual(len(turni), 14)
        lunedi = timezone.localdate() - timedelta(days=timezone.localdate().weekday())
        self.assertEqual(turni[0].data, lunedi)
        self.assertEqual({t.sigla for t in turni}, {'1T', '2T', 'CT', 'N', 'G', 'R', 'F'})
        # una seconda volta non fa nulla
        self.client.post(url, {'dipendente': dip.pk, 'azione': 'turni_prova'})
        self.assertEqual(dip.turni.count(), 14)
        self.assertNotContains(self.client.get(url), 'Genera turni di prova')


class SettimanaTuttiTest(TestCase):
    def setUp(self):
        importa.salva_settimana(importa.leggi_file(file_turni(RIGHE), LUNEDI))
        d = Dipendente.objects.get(nome='DRAME BOUBA')
        d.utente = get_user_model().objects.create_user('drame', password='123456')
        d.save()
        self.client.login(username='drame', password='123456')
        adesso = timezone.make_aware(datetime.combine(LUNEDI, datetime.min.time()) + timedelta(hours=8))
        patcher = mock.patch('django.utils.timezone.now', return_value=adesso)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_griglia_per_impianto_e_celle_cambiabili(self):
        r = self.client.get(reverse('turni:settimana_tutti'))
        self.assertEqual(r.status_code, 200)
        gruppi = dict(r.context['gruppi'])
        # Drame e Spano lavorano soprattutto all'Impianto 1, Crimi soprattutto da palista
        self.assertEqual({x['nome'] for x in gruppi['Impianto 1']}, {'Spano’ A. M.', 'Drame B.'})
        self.assertIn('Crimi B.', {x['nome'] for x in gruppi['Palista']})
        spano = next(x for x in gruppi['Impianto 1'] if x['nome'].startswith('Spano'))
        # lunedì (oggi) non si cambia; martedì Spano riposa e Drame lavora: si può chiedere
        self.assertFalse(spano['celle'][0]['cambiabile'])
        self.assertTrue(spano['celle'][1]['cambiabile'])
        # giovedì hanno lo stesso turno: niente da cambiare
        self.assertFalse(spano['celle'][3]['cambiabile'])
        io = next(x for x in gruppi['Impianto 1'] if x['io'])
        self.assertFalse(any(c['cambiabile'] for c in io['celle']))
        self.assertContains(r, f"collega={spano['dip'].pk}")


class PostazioniOrariTest(TestCase):
    def test_nuove_postazioni_e_orari(self):
        casi = {
            '1T-Metalli': (Fascia.MATTINA, Postazione.METALLI),
            'G - Sabbia': (Fascia.GIORNALIERO, Postazione.SABBIA),
            '2T/Piazzale': (Fascia.POMERIGGIO, Postazione.PIAZZALE),
            'G-Autista': (Fascia.GIORNALIERO, Postazione.AUTISTA),
            'N/Manutentore': (Fascia.NOTTE, Postazione.MANUT),
            'G-Ufficio': (Fascia.GIORNALIERO, Postazione.UFFICIO),
            'Notturno': (Fascia.NOTTE, ''),
        }
        for codice, atteso in casi.items():
            c = importa.interpreta_codice(codice)
            self.assertEqual((c.fascia, c.postazione), atteso, codice)
        t = Turno(fascia=Fascia.MATTINA, postazione=Postazione.METALLI)
        self.assertEqual((t.orario, t.orario_pulizie, t.sigla_postazione), ('06:00–14:00', '13:30–14:00', 'MET.'))
        self.assertEqual(Turno(fascia=Fascia.GIORNALIERO).orario_pulizie, '17:00–17:30')
        self.assertEqual(Turno(fascia=Fascia.NOTTE).descrizione_fascia, 'Notturno')
        self.assertEqual(Turno(fascia=Fascia.RIPOSO).orario, '')
