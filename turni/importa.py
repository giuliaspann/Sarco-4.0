"""Lettura del file turni settimanale (quello condiviso su WhatsApp).

Il file contiene un foglio per settimana ("Sett.1" ... "Sett.53") numerato come la
settimana ISO: "Sett.42" del 2026 va dal 12/10 al 18/10. In ogni foglio la tabella
inizia dalla riga con "DIPENDENTI" in colonna B, con i giorni da lunedì a domenica
nelle colonne C-I. A destra (K-M) ci sono le note su ferie e richieste, in fondo
il riquadro "APPRENDISTI E TUTOR".
"""

import re
from dataclasses import dataclass, field
from datetime import date, timedelta

import openpyxl

from .models import Fascia, Postazione

POSTAZIONI = [
    (r'IMP\.?\s*1', Postazione.IMP1),
    (r'IMP\.?\s*2', Postazione.IMP2),
    (r'ANALISI', Postazione.ANALISI),
    (r'SUPERBI', Postazione.SUPERBI),
    (r'MULINO', Postazione.MULINO),
    (r'BARROTTI', Postazione.BARROTTI),
    (r'PALISTA', Postazione.PALISTA),
]

ASSENZE = {
    'RIP': Fascia.RIPOSO,
    'RIPOSO': Fascia.RIPOSO,
    'FERIE': Fascia.FERIE,
    'FESTA': Fascia.FESTA,
    'C.P.': Fascia.PERMESSO,
    'CP': Fascia.PERMESSO,
}


@dataclass
class CodiceTurno:
    codice: str
    fascia: str
    postazione: str = ''
    capoturno: bool = False


def interpreta_codice(testo):
    """'1T- Imp. 1 ' -> 1° turno Impianto 1; 'CAPO T. Imp. 1 N' -> capoturno notte Imp. 1; ..."""
    codice = ' '.join(str(testo or '').split())
    t = codice.upper()
    if not t:
        return CodiceTurno('', Fascia.RIPOSO)
    if t in ASSENZE:
        return CodiceTurno(codice, ASSENZE[t])

    capoturno = bool(re.match(r'CAPO\s*T', t))
    resto = re.sub(r'^CAPO\s*T\.?', '', t).strip(' -/')

    postazione = ''
    for modello, valore in POSTAZIONI:
        if re.search(modello, resto):
            postazione = valore
            resto = re.sub(modello, ' ', resto)
            break

    # quello che resta indica la fascia: 1T, 2T, N, G
    token = set(re.split(r'[\s\-/]+', resto)) - {''}
    if '1T' in token:
        fascia = Fascia.MATTINA
    elif '2T' in token:
        fascia = Fascia.POMERIGGIO
    elif 'N' in token:
        fascia = Fascia.NOTTE
    elif 'G' in token:
        fascia = Fascia.GIORNALIERO
    else:
        fascia = Fascia.ALTRO  # es. "CAPO T. - Imp. 1": la fascia non è scritta
    return CodiceTurno(codice, fascia, postazione, capoturno)


@dataclass
class RigaDipendente:
    nome: str
    turni: list  # 7 CodiceTurno, da lunedì a domenica


@dataclass
class SettimanaLetta:
    lunedi: date
    foglio: str
    righe: list = field(default_factory=list)
    note: list = field(default_factory=list)  # (riferimento, dipendente, testo)
    tutor: dict = field(default_factory=dict)  # apprendista -> tutor
    avvisi: list = field(default_factory=list)


def lunedi_da_nome_file(nome):
    m = re.search(r'(\d{1,2})[-.](\d{1,2})[-.](\d{4})', nome or '')
    if not m:
        return None
    g, me, a = map(int, m.groups())
    d = date(a, me, g)
    return d - timedelta(days=d.weekday())


def pulisci(v):
    return ' '.join(str(v).split()) if v is not None else ''


def leggi_file(file, lunedi):
    wb = openpyxl.load_workbook(file, data_only=True)
    anno, settimana, _ = lunedi.isocalendar()
    nome_foglio = f'Sett.{settimana}'
    if nome_foglio in wb.sheetnames:
        ws = wb[nome_foglio]
        avvisi = []
    else:
        ws = wb.worksheets[-1]
        avvisi = [f'Foglio "{nome_foglio}" non trovato: uso l\'ultimo foglio "{ws.title}".']

    letta = SettimanaLetta(lunedi=lunedi, foglio=ws.title, avvisi=avvisi)

    intestazione = None
    for r in range(1, 15):
        if pulisci(ws.cell(r, 2).value).upper() == 'DIPENDENTI':
            intestazione = r
            break
    if intestazione is None:
        raise ValueError(f'Nel foglio "{ws.title}" non trovo la riga "DIPENDENTI" in colonna B.')

    r = intestazione + 1
    while True:
        nome = pulisci(ws.cell(r, 2).value)
        if not nome:
            break
        turni = [interpreta_codice(ws.cell(r, c).value) for c in range(3, 10)]
        letta.righe.append(RigaDipendente(nome.upper(), turni))
        r += 1

    # note a destra (K = settimana n., L = dipendente, M = testo), anche sulla riga d'intestazione
    for r in range(intestazione, ws.max_row + 1):
        rif, dip, testo = (pulisci(ws.cell(r, c).value) for c in (11, 12, 13))
        if dip.upper() == 'DIPENDENTI' or testo.upper() == 'RICHIESTA':
            continue
        if testo or dip:
            letta.note.append((rif, dip, testo))

    # riquadro apprendisti: "NOME (TUTOR1 E TUTOR2)" anche spezzato su due celle
    for r in range(1, ws.max_row + 1):
        for c in range(1, 17):
            if 'APPRENDISTI' in pulisci(ws.cell(r, c).value).upper():
                rr = r + 1
                while True:
                    riga = ' '.join(filter(None, (pulisci(ws.cell(rr, cc).value) for cc in range(c, c + 4))))
                    m = re.match(r'(.+?)\s*\((.+)\)', riga)
                    if not m:
                        break
                    letta.tutor[m.group(1).strip().upper()] = m.group(2).strip()
                    rr += 1
                break
    return letta


# ---------------------------------------------------------------- confronto nomi

def _token(nome):
    nome = re.sub(r"[’'`.]", ' ', (nome or '').upper())
    return [t for t in nome.split() if len(t) > 1]


def stesso_nome(a, b):
    """I nomi nel file non sono sempre scritti uguali: 'SPANO MAURO' / "SPANO' A. MAURO",
    'CARLO SAMMARTANO' / 'SAMMARTANO CARLO', 'DRAME BUOBA' / 'DRAME BOUBA'."""
    import difflib

    ta, tb = _token(a), _token(b)
    if not ta or not tb:
        return False
    sa, sb = set(ta), set(tb)
    if len(sa & sb) >= 2 and (sa <= sb or sb <= sa):
        return True
    return difflib.SequenceMatcher(None, ' '.join(sorted(ta)), ' '.join(sorted(tb))).ratio() >= 0.85


def trova_dipendente(nome, dipendenti):
    for d in dipendenti:
        if d.nome == nome.upper():
            return d
    for d in dipendenti:
        if stesso_nome(d.nome, nome):
            return d
    return None


# ---------------------------------------------------------------- salvataggio

def in_sessione(letta):
    """La settimana letta, in forma serializzabile per l'anteprima prima della conferma."""
    return {
        'lunedi': letta.lunedi.isoformat(),
        'foglio': letta.foglio,
        'avvisi': letta.avvisi,
        'righe': [[r.nome, [[t.codice, t.fascia, t.postazione, t.capoturno] for t in r.turni]] for r in letta.righe],
        'note': letta.note,
        'tutor': letta.tutor,
    }


def da_sessione(dati):
    letta = SettimanaLetta(
        lunedi=date.fromisoformat(dati['lunedi']),
        foglio=dati['foglio'],
        note=[tuple(n) for n in dati['note']],
        tutor=dati['tutor'],
        avvisi=dati['avvisi'],
    )
    letta.righe = [RigaDipendente(nome, [CodiceTurno(*t) for t in turni]) for nome, turni in dati['righe']]
    return letta


def salva_settimana(letta, utente=None, nome_file=''):
    """Crea o aggiorna la settimana. Ricaricare il file corretto a settimana iniziata
    aggiorna i turni senza perdere le pulizie già registrate.

    I turni modificati nell'app (cambi e ferie approvati) restano come sono: la terza
    lista restituita dice quali differivano dal file, per avvisare l'ufficio."""
    from django.db import transaction

    from .models import Dipendente, NotaSettimana, Settimana, Turno
    from .richieste import ferie_approvate, segna_ferie_sui_turni

    with transaction.atomic():
        settimana, _ = Settimana.objects.update_or_create(
            lunedi=letta.lunedi,
            defaults={'file_originale': nome_file, 'foglio': letta.foglio, 'caricata_da': utente},
        )
        dipendenti = list(Dipendente.objects.all())
        nuovi, visti, mantenuti = [], set(), []
        for riga in letta.righe:
            dip = trova_dipendente(riga.nome, dipendenti)
            if dip is None:
                dip = Dipendente.objects.create(nome=riga.nome)
                dipendenti.append(dip)
                nuovi.append(dip.nome)
            visti.add(dip.pk)
            esistenti = {t.data: t for t in Turno.objects.filter(dipendente=dip, data__in=settimana.giorni)}
            for giorno, t in zip(settimana.giorni, riga.turni):
                vecchio = esistenti.get(giorno)
                if vecchio and vecchio.nota_modifica:
                    # cambio turno o ferie approvati nell'app: hanno la precedenza sul file
                    if (vecchio.fascia, vecchio.postazione, vecchio.capoturno) != (t.fascia, t.postazione, t.capoturno):
                        mantenuti.append(f'{dip.nome_breve} {giorno:%d/%m}: {vecchio.descrizione} ({vecchio.nota_modifica})')
                    continue
                nuovo, _ = Turno.objects.update_or_create(
                    dipendente=dip,
                    data=giorno,
                    defaults={
                        'settimana': settimana,
                        'codice': t.codice,
                        'fascia': t.fascia,
                        'postazione': t.postazione,
                        'capoturno': t.capoturno,
                    },
                )
                ferie = ferie_approvate(dip, giorno)
                if ferie and nuovo.fascia != ferie.fascia:
                    segna_ferie_sui_turni(ferie, [nuovo])
                    mantenuti.append(f'{dip.nome_breve} {giorno:%d/%m}: {ferie.get_tipo_display()} approvate')
        # chi non è più nel file perde i turni della settimana, a meno che abbia già pulizie registrate
        settimana.turni.exclude(dipendente__in=visti).filter(pulizie__isnull=True, nota_modifica='').delete()

        settimana.note.all().delete()
        NotaSettimana.objects.bulk_create(
            NotaSettimana(settimana=settimana, riferimento=rif, dipendente=dip, testo=testo or '-', ordine=i)
            for i, (rif, dip, testo) in enumerate(letta.note)
        )
        for apprendista, tutor in letta.tutor.items():
            dip = trova_dipendente(apprendista, dipendenti)
            if dip and dip.tutor != tutor:
                dip.tutor = tutor
                dip.save(update_fields=['tutor'])
    return settimana, nuovi, mantenuti
